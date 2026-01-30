from pptx import Presentation
import os

# Load the PowerPoint template
prs = Presentation('/home/user/Cardamyst/CARDAMYST Branded Powerpoint Template.pptx')

print(f"Total slides: {len(prs.slides)}")
print("\nSearching for images (logo)...")

# Create directory for extracted images
os.makedirs('/home/user/Cardamyst/extracted_assets', exist_ok=True)

# Extract all images from all slides
image_count = 0
for slide_num, slide in enumerate(prs.slides, 1):
    print(f"\nSlide {slide_num}:")
    for shape in slide.shapes:
        if shape.shape_type == 13:  # Picture type
            image = shape.image
            image_bytes = image.blob
            image_ext = image.ext
            image_filename = f'/home/user/Cardamyst/extracted_assets/slide{slide_num}_image{image_count}.{image_ext}'

            with open(image_filename, 'wb') as f:
                f.write(image_bytes)

            print(f"  - Extracted image: {image_filename} ({len(image_bytes)} bytes)")
            image_count += 1

print(f"\n\nTotal images extracted: {image_count}")
print("Images saved to: /home/user/Cardamyst/extracted_assets/")
