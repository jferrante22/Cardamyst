import openpyxl
import json
from datetime import datetime

# Load the Excel file
file_path = '/home/user/Cardamyst/Cardamyst_Coverage_Tracker Jan 2026.xlsx'
wb = openpyxl.load_workbook(file_path, data_only=True)

# Extract Executive Summary data
exec_summary = wb['Executive Summary']
monthly_tracking = wb['Monthly Tracking']
charts = wb['Charts']

# Parse coverage snapshot
dashboard_data = {
    "report_date": "January 2026",
    "last_updated": datetime.now().isoformat(),
    "coverage_snapshot": {
        "commercial": {
            "total_lives": 186512403,
            "covered_lives": 6816075,
            "coverage_percent": 3.65,
            "notes": "Employer, PBM, HIX, FEHBP, VA, TRICARE"
        },
        "federal_programs": {
            "total_lives": 15711264,
            "covered_lives": 5749433,
            "coverage_percent": 36.59,
            "notes": "TRICARE, VA (subset of Commercial)"
        },
        "medicare": {
            "total_lives": 53706458,
            "covered_lives": 0,
            "coverage_percent": 0,
            "notes": "Part D, MA, SNP, EGWP, PACE"
        },
        "medicaid": {
            "total_lives": 76473937,
            "covered_lives": 0,
            "coverage_percent": 0,
            "notes": "State, Managed, CHIP"
        },
        "total": {
            "total_lives": 316692798,
            "covered_lives": 6816075,
            "coverage_percent": 2.15
        }
    },
    "top_plans": [],
    "monthly_forecast": {
        "commercial": {
            "months": [],
            "forecast": [],
            "actual": []
        }
    }
}

# Extract top plans
for row in exec_summary.iter_rows(min_row=16, max_row=25, values_only=True):
    if row[0] and row[0] != 'Plan Name':
        plan = {
            "name": row[0],
            "type": row[1] if row[1] else "",
            "covered_lives": int(row[2]) if row[2] else 0,
            "status": row[3] if row[3] else ""
        }
        dashboard_data["top_plans"].append(plan)

# Extract monthly forecast data from Charts sheet
for row in charts.iter_rows(min_row=4, max_row=18, values_only=True):
    if row[0]:  # Month name
        dashboard_data["monthly_forecast"]["commercial"]["months"].append(row[0])
        dashboard_data["monthly_forecast"]["commercial"]["forecast"].append(
            float(row[1]) * 100 if row[1] is not None else None
        )
        dashboard_data["monthly_forecast"]["commercial"]["actual"].append(
            float(row[2]) * 100 if row[2] is not None else None
        )

# Save to JSON
with open('/home/user/Cardamyst/dashboard_data.json', 'w') as f:
    json.dump(dashboard_data, f, indent=2)

print("Dashboard data extracted successfully!")
print(json.dumps(dashboard_data, indent=2))

wb.close()
