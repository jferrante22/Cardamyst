import openpyxl
import json
import os

# Load the Excel file
file_path = '/home/user/Cardamyst/Cardamyst_Coverage_Tracker Jan 2026.xlsx'
wb = openpyxl.load_workbook(file_path, data_only=True)

# Print all sheet names
print("Available sheets:", wb.sheetnames)
print("\n" + "="*80 + "\n")

# Extract data from each sheet
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"Sheet: {sheet_name}")
    print(f"Max Row: {ws.max_row}, Max Column: {ws.max_column}")
    print("\nFirst 20 rows:")

    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i <= 20:
            print(f"Row {i}: {row}")

    print("\n" + "="*80 + "\n")

wb.close()
