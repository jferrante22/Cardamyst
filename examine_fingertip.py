import openpyxl

# Load the January Fingertip file
file_path = '/home/user/Cardamyst/January Fingertip 01.30.xlsx'
wb = openpyxl.load_workbook(file_path, data_only=True)

# Print all sheet names
print("Available sheets:", wb.sheetnames)
print("\n" + "="*80 + "\n")

# Examine each sheet
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"Sheet: {sheet_name}")
    print(f"Max Row: {ws.max_row}, Max Column: {ws.max_column}")
    print("\nFirst 30 rows:")

    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i <= 30:
            print(f"Row {i}: {row}")

    print("\n" + "="*80 + "\n")

wb.close()
